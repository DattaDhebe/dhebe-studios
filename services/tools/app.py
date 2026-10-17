import io
import os
from pathlib import Path
import shutil
import subprocess
import tempfile
import threading
import zipfile
import secrets
import time
from concurrent.futures import ThreadPoolExecutor

from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import Response
from starlette.datastructures import UploadFile as MemoryUpload
from pypdf import PdfReader, PdfWriter
from early_access import router as early_access_router

MAX_BYTES = 25 * 1024 * 1024
MAX_PAGES = 150
SLOTS = threading.BoundedSemaphore(2)
POOL = ThreadPoolExecutor(max_workers=2)
JOBS = {}
JOBS_LOCK = threading.Lock()


def expire_jobs():
    while True:
        time.sleep(30)
        now = time.monotonic()
        with JOBS_LOCK:
            for key in list(JOBS):
                started, future = JOBS[key]
                if future.done() and now - started > 300:
                    del JOBS[key]


threading.Thread(target=expire_jobs, daemon=True).start()
EXTENSIONS = {
    'word-to-pdf': {'.doc', '.docx', '.odt'},
    'powerpoint-to-pdf': {'.ppt', '.pptx', '.odp'},
    'excel-to-pdf': {'.xls', '.xlsx', '.ods', '.csv'},
    'html-to-pdf': {'.html', '.htm'},
    'background-remover': {'.png', '.jpg', '.jpeg', '.webp'},
    **{name: {'.pdf'} for name in ['pdf-to-word', 'pdf-to-powerpoint', 'pdf-to-excel', 'pdf-to-pdfa', 'protect', 'ocr', 'repair']},
}
MIME = {'.pdf': 'application/pdf', '.docx': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document', '.pptx': 'application/vnd.openxmlformats-officedocument.presentationml.presentation', '.xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', '.png': 'image/png'}
app = FastAPI(docs_url=None, redoc_url=None)
app.include_router(early_access_router)
origins = [origin.strip() for origin in os.getenv('ALLOWED_ORIGINS', '').split(',') if origin.strip()]
if origins:
    app.add_middleware(CORSMiddleware, allow_origins=origins, allow_methods=['POST', 'GET'], allow_headers=['Content-Type'], expose_headers=['Content-Disposition'])


def command(args, timeout=180, allowed=(0,)):
    try:
        result = subprocess.run(args, capture_output=True, timeout=timeout, check=False)
    except subprocess.TimeoutExpired:
        raise HTTPException(504, 'Conversion timed out. Try a smaller document.') from None
    if result.returncode not in allowed:
        raise HTTPException(422, 'This document could not be converted. Check that it opens correctly and is not password-protected.')
    return result


def validate_archive(source):
    if source.suffix in {'.docx', '.pptx', '.xlsx', '.odt', '.odp', '.ods'}:
        try:
            with zipfile.ZipFile(source) as archive:
                if sum(item.file_size for item in archive.infolist()) > 200 * 1024 * 1024:
                    raise HTTPException(413, 'The document expands beyond the processing limit.')
        except zipfile.BadZipFile:
            raise HTTPException(422, 'The uploaded document is not a valid Office file.') from None


def read_pdf(source):
    try:
        reader = PdfReader(source)
        if reader.is_encrypted:
            raise HTTPException(422, 'Unlock this PDF before converting it.')
        if not 0 < len(reader.pages) <= MAX_PAGES:
            raise HTTPException(413, f'Choose a PDF with 1 to {MAX_PAGES} pages.')
        return reader
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(422, 'This PDF cannot be read. Try Repair PDF first.') from None


def convert(source, tool, password=''):
    work = source.parent
    if tool in {'word-to-pdf', 'powerpoint-to-pdf', 'excel-to-pdf', 'html-to-pdf'}:
        validate_archive(source)
        profile = (work / 'lo-profile').as_uri()
        command(['libreoffice', f'-env:UserInstallation={profile}', '--headless', '--convert-to', 'pdf', '--outdir', str(work), str(source)])
        output = source.with_suffix('.pdf')
        if not output.is_file():
            raise HTTPException(422, 'The Office document could not be converted.')
        read_pdf(output)
        return output
    if tool == 'background-remover':
        from PIL import Image
        from rembg import new_session, remove
        Image.MAX_IMAGE_PIXELS = 16_000_000
        with Image.open(source) as image:
            if image.width * image.height > 16_000_000:
                raise HTTPException(413, 'Choose an image smaller than 16 megapixels.')
            output = work / 'result.png'
            remove(image, session=new_session('u2netp')).save(output)
        return output
    if tool == 'repair':
        output = work / 'result.pdf'
        command(['qpdf', str(source), str(output)], allowed=(0, 3))
        read_pdf(output)
        return output
    reader = read_pdf(source)
    if tool == 'protect':
        if len(password) < 8 or len(password) > 128:
            raise HTTPException(400, 'Use a password with 8 to 128 characters.')
        writer = PdfWriter(clone_from=reader)
        writer.encrypt(password, algorithm='AES-256')
        output = work / 'result.pdf'
        writer.write(output)
        return output
    if tool == 'pdf-to-word':
        from docx import Document
        document = Document()
        found = False
        for index, page in enumerate(reader.pages):
            text = page.extract_text() or ''
            found = found or bool(text.strip())
            if index:
                document.add_page_break()
            for line in text.splitlines():
                document.add_paragraph(line)
        if not found:
            raise HTTPException(422, 'No selectable text found. Run OCR before converting this PDF.')
        output = work / 'result.docx'
        document.save(output)
        return output
    if tool == 'pdf-to-powerpoint':
        import pypdfium2 as pdfium
        from pptx import Presentation
        from pptx.util import Inches
        presentation = Presentation()
        first = reader.pages[0].mediabox
        presentation.slide_width = Inches(float(first.width) / 72)
        presentation.slide_height = Inches(float(first.height) / 72)
        with pdfium.PdfDocument(source) as pdf:
            for index in range(len(pdf)):
                page = pdf[index]
                bitmap = page.render(scale=min(1.5, 2048 / max(page.get_size())))
                image = bitmap.to_pil()
                data = io.BytesIO()
                image.save(data, format='PNG')
                data.seek(0)
                slide = presentation.slides.add_slide(presentation.slide_layouts[6])
                ratio = min(presentation.slide_width / image.width, presentation.slide_height / image.height)
                width, height = int(image.width * ratio), int(image.height * ratio)
                slide.shapes.add_picture(data, (presentation.slide_width - width) // 2, (presentation.slide_height - height) // 2, width=width, height=height)
                image.close()
                bitmap.close()
                page.close()
        output = work / 'result.pptx'
        presentation.save(output)
        return output
    if tool == 'pdf-to-excel':
        import pdfplumber
        from openpyxl import Workbook
        workbook = Workbook()
        workbook.remove(workbook.active)
        found = False
        with pdfplumber.open(source) as pdf:
            for index, page in enumerate(pdf.pages):
                sheet = workbook.create_sheet(f'Page {index + 1}')
                tables = page.extract_tables()
                rows = [row for table in tables for row in table] if tables else [[line] for line in (page.extract_text() or '').splitlines()]
                for row in rows:
                    values = [str(value or '')[:32767] for value in row]
                    found = found or any(value.strip() for value in values)
                    sheet.append(values)
                    # Untrusted PDF text must stay text, even when it starts with '='.
                    for cell in sheet[sheet.max_row]:
                        cell.data_type = 's'
        if not found:
            raise HTTPException(422, 'No tables or selectable text found. Run OCR first.')
        output = work / 'result.xlsx'
        workbook.save(output)
        return output
    if tool == 'ocr':
        output = work / 'result.pdf'
        command(['ocrmypdf', '--skip-text', '--jobs', '1', '--tesseract-timeout', '60', '--output-type', 'pdf', str(source), str(output)], timeout=240)
        return output
    if tool == 'pdf-to-pdfa':
        output = work / 'result.pdf'
        command(['ocrmypdf', '--skip-text', '--jobs', '1', '--output-type', 'pdfa-2', str(source), str(output)], timeout=240)
        converted = PdfReader(output)
        metadata = converted.xmp_metadata
        if not metadata or metadata.pdfaid_part != '2' or metadata.pdfaid_conformance != 'B':
            raise HTTPException(422, 'The document could not be converted to PDF/A-2b.')
        return output
    raise HTTPException(404, 'Unknown conversion tool.')


@app.get('/health')
def health():
    return {'status': 'ok', 'engines': {name: bool(shutil.which(name)) for name in ['libreoffice', 'qpdf', 'ocrmypdf']}}


@app.post('/jobs/{tool}', status_code=202)
def start_job(tool: str, file: UploadFile = File(...), password: str = Form('')):
    if tool not in EXTENSIONS:
        raise HTTPException(404, 'Unknown conversion tool.')
    if Path(file.filename or '').suffix.lower() not in EXTENSIONS[tool]:
        raise HTTPException(415, 'This file type is not supported by the selected tool.')
    content = file.file.read(MAX_BYTES + 1)
    file.file.close()
    if len(content) > MAX_BYTES:
        raise HTTPException(413, 'Choose a file smaller than 25 MB.')
    if not content:
        raise HTTPException(400, 'The uploaded file is empty.')
    with JOBS_LOCK:
        now = time.monotonic()
        for key in list(JOBS):
            started, future = JOBS[key]
            if future.done() and now - started > 300:
                del JOBS[key]
        if len(JOBS) >= 4:
            raise HTTPException(429, 'All conversion slots are busy. Please try again shortly.')
        token = secrets.token_urlsafe(32)
        upload = MemoryUpload(file=io.BytesIO(content), filename=file.filename)
        JOBS[token] = (now, POOL.submit(process, tool, upload, password))
    return {'job': token}


@app.get('/jobs/{token}')
def job_result(token: str):
    with JOBS_LOCK:
        entry = JOBS.get(token)
        if not entry:
            raise HTTPException(404, 'The download expired. Convert the file again.')
        _, future = entry
        if not future.done():
            return Response(status_code=202, headers={'Cache-Control': 'no-store'})
        del JOBS[token]
    return future.result()


@app.post('/convert/{tool}')
def process(tool: str, file: UploadFile = File(...), password: str = Form('')):
    if tool not in EXTENSIONS:
        raise HTTPException(404, 'Unknown conversion tool.')
    extension = Path(file.filename or '').suffix.lower()
    if extension not in EXTENSIONS[tool]:
        raise HTTPException(415, 'This file type is not supported by the selected tool.')
    if not SLOTS.acquire(blocking=False):
        raise HTTPException(429, 'All conversion slots are busy. Please try again shortly.')
    try:
        with tempfile.TemporaryDirectory(prefix='dhebe-convert-') as directory:
            source = Path(directory) / f'input{extension}'
            size = 0
            with source.open('wb') as target:
                while chunk := file.file.read(1024 * 1024):
                    size += len(chunk)
                    if size > MAX_BYTES:
                        raise HTTPException(413, 'Choose a file smaller than 25 MB.')
                    target.write(chunk)
            if not size:
                raise HTTPException(400, 'The uploaded file is empty.')
            output = convert(source, tool, password)
            if not output.is_file() or output.stat().st_size == 0:
                raise HTTPException(422, 'The conversion did not produce a file.')
            if output.stat().st_size > 100 * 1024 * 1024:
                raise HTTPException(413, 'The converted file exceeds the download limit.')
            content = output.read_bytes()
            return Response(content, media_type=MIME[output.suffix], headers={'Content-Disposition': f'attachment; filename="{tool}{output.suffix}"', 'Cache-Control': 'no-store', 'X-Content-Type-Options': 'nosniff'})
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(422, 'The document could not be processed. Check the file and try again.') from None
    finally:
        file.file.close()
        SLOTS.release()
