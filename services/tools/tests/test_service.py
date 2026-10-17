import io
from pathlib import Path
import tempfile
import unittest
import time
from fastapi.testclient import TestClient
from docx import Document
from pptx import Presentation
from openpyxl import Workbook, load_workbook
from pypdf import PdfReader
from PIL import Image, ImageDraw, ImageFont
from app import app

client = TestClient(app)


def post(tool, name, content, **options):
    response = client.post(f'/convert/{tool}', files={'file': (name, content)}, data=options)
    if response.status_code != 200:
        raise AssertionError(f'{tool}: {response.status_code}: {response.text}')
    assert response.headers['cache-control'] == 'no-store'
    return response.content


class ConversionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        document = Document()
        document.add_heading('Conversion regression', 0)
        document.add_paragraph('Invoice total 125.50')
        document.add_paragraph('=HYPERLINK("https://example.com", "example")')
        data = io.BytesIO()
        document.save(data)
        cls.word = data.getvalue()
        cls.pdf = post('word-to-pdf', 'sample.docx', cls.word)

    def test_office_inputs(self):
        self.assertIn('Invoice total', PdfReader(io.BytesIO(self.pdf)).pages[0].extract_text())
        presentation = Presentation()
        presentation.slides.add_slide(presentation.slide_layouts[1]).shapes.title.text = 'Presentation test'
        data = io.BytesIO()
        presentation.save(data)
        self.assertTrue(post('powerpoint-to-pdf', 'sample.pptx', data.getvalue()).startswith(b'%PDF'))
        workbook = Workbook()
        workbook.active.append(['Name', 'Total'])
        workbook.active.append(['Invoice', 125.5])
        data = io.BytesIO()
        workbook.save(data)
        self.assertTrue(post('excel-to-pdf', 'sample.xlsx', data.getvalue()).startswith(b'%PDF'))
        self.assertTrue(post('html-to-pdf', 'sample.html', b'<h1>HTML heading</h1><p>Text</p>').startswith(b'%PDF'))

    def test_real_office_outputs(self):
        word = Document(io.BytesIO(post('pdf-to-word', 'sample.pdf', self.pdf)))
        self.assertIn('Invoice', '\n'.join(p.text for p in word.paragraphs))
        slides = Presentation(io.BytesIO(post('pdf-to-powerpoint', 'sample.pdf', self.pdf)))
        self.assertEqual(len(slides.slides), len(PdfReader(io.BytesIO(self.pdf)).pages))
        workbook = load_workbook(io.BytesIO(post('pdf-to-excel', 'sample.pdf', self.pdf)))
        cells = [cell for row in workbook.active for cell in row]
        self.assertTrue(any('Invoice' in str(cell.value) for cell in cells))
        self.assertFalse(any(cell.data_type == 'f' for cell in cells))

    def test_password_encryption(self):
        encrypted = PdfReader(io.BytesIO(post('protect', 'sample.pdf', self.pdf, password='test-password-123')))
        self.assertTrue(encrypted.is_encrypted)
        self.assertFalse(encrypted.decrypt('wrong-password'))
        self.assertTrue(encrypted.decrypt('test-password-123'))
        self.assertEqual(len(encrypted.pages), 1)

    def test_pdfa(self):
        reader = PdfReader(io.BytesIO(post('pdf-to-pdfa', 'sample.pdf', self.pdf)))
        self.assertEqual(str(reader.xmp_metadata.pdfaid_part), '2')
        self.assertEqual(reader.xmp_metadata.pdfaid_conformance, 'B')
        self.assertIn('/OutputIntents', reader.trailer['/Root'])

    def test_ocr_all_pages(self):
        font = ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf', 65)
        images = []
        for text in ['Invoice total 125', 'Second page receipt']:
            image = Image.new('RGB', (1500, 700), 'white')
            ImageDraw.Draw(image).text((100, 200), text, fill='black', font=font)
            images.append(image)
        data = io.BytesIO()
        images[0].save(data, format='PDF', save_all=True, append_images=images[1:], resolution=150)
        result = PdfReader(io.BytesIO(post('ocr', 'scan.pdf', data.getvalue())))
        self.assertEqual(len(result.pages), 2)
        self.assertIn('Invoice', result.pages[0].extract_text())
        self.assertIn('Second', result.pages[1].extract_text())

    def test_repair_and_cleanup(self):
        before = set(Path(tempfile.gettempdir()).glob('dhebe-convert-*'))
        self.assertTrue(post('repair', '../../sample.pdf', self.pdf).startswith(b'%PDF'))
        self.assertEqual(client.post('/convert/pdf-to-word', files={'file': ('bad.pdf', b'bad')}).status_code, 422)
        self.assertEqual(set(Path(tempfile.gettempdir()).glob('dhebe-convert-*')), before)

    def test_validation(self):
        for tool, name, data, status in [('unknown', 'a.pdf', self.pdf, 404), ('pdf-to-word', 'a.exe', self.pdf, 415), ('pdf-to-word', 'a.pdf', b'', 400), ('protect', 'a.pdf', self.pdf, 400)]:
            self.assertEqual(client.post(f'/convert/{tool}', files={'file': (name, data)}).status_code, status)

    def test_async_jobs(self):
        response = client.post('/jobs/pdf-to-word', files={'file': ('sample.pdf', self.pdf)})
        self.assertEqual(response.status_code, 202)
        token = response.json()['job']
        for _ in range(100):
            result = client.get('/jobs/' + token)
            if result.status_code != 202:
                break
            time.sleep(0.1)
        self.assertEqual(result.status_code, 200)
        self.assertTrue(result.content.startswith(b'PK'))
        self.assertEqual(client.get('/jobs/' + token).status_code, 404)

    def test_background_removal(self):
        image = Image.new('RGB', (150, 150), 'white')
        ImageDraw.Draw(image).ellipse((25, 20, 125, 135), fill='red')
        data = io.BytesIO()
        image.save(data, format='PNG')
        output = Image.open(io.BytesIO(post('background-remover', 'image.png', data.getvalue())))
        self.assertEqual(output.mode, 'RGBA')
        self.assertEqual(output.size, (150, 150))


if __name__ == '__main__':
    unittest.main(verbosity=2)
